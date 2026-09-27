from __future__ import annotations

import io
from typing import Any


def build_docx(payload: dict[str, Any]) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as error:
        raise RuntimeError("La dependencia python-docx no está instalada.") from error

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    title = document.add_heading(payload["project"]["name"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Dominio: {payload['project'].get('domain') or 'No especificado'}")
    if payload["project"].get("description"):
        document.add_paragraph(payload["project"]["description"])

    profile = payload.get("profile") or {}
    if profile:
        document.add_heading("Definición confirmada", level=1)
        labels = {
            "project_goal": "Objetivo",
            "problem": "Problema",
            "actors": "Actores",
            "scope": "Alcance",
            "out_of_scope": "Fuera de alcance",
            "business_and_legal_constraints": "Reglas y restricciones",
            "quality_expectations": "Expectativas de calidad",
            "known_conflicts": "Conflictos conocidos",
        }
        for key, label in labels.items():
            if profile.get(key):
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{label}: ").bold = True
                paragraph.add_run(str(profile[key]))

    for artifact_type, heading in (("RF", "Requisitos funcionales"), ("RNF", "Requisitos no funcionales"), ("HU", "Historias de usuario")):
        artifacts = [item for item in payload.get("artifacts", []) if item["artifact_type"] == artifact_type]
        document.add_heading(heading, level=1)
        if not artifacts:
            document.add_paragraph("No se generaron artefactos de este tipo.")
            continue
        for item in artifacts:
            _add_artifact_table(document, item, artifact_type)
            document.add_paragraph()

    fragments = {item["fragment_id"]: item for item in payload.get("fragments", [])}
    document.add_heading("Matriz de trazabilidad", level=1)
    matrix = document.add_table(rows=1, cols=6)
    matrix.style = "Table Grid"
    for cell, label in zip(
        matrix.rows[0].cells,
        ("Artefacto", "Tipo", "Fragmento", "Fuente", "Extracto de evidencia", "Estado"),
    ):
        cell.text = label
    for item in payload.get("artifacts", []):
        citations = item.get("source_fragments") or [""]
        for citation in citations:
            evidence = fragments.get(citation, {})
            cells = matrix.add_row().cells
            cells[0].text = str(item["artifact_key"])
            cells[1].text = str(item["artifact_type"])
            cells[2].text = str(citation or "Sin cita")
            cells[3].text = str(evidence.get("source_file") or "No disponible")
            excerpt = " ".join(str(evidence.get("text") or "").split())
            cells[4].text = excerpt[:280] + ("…" if len(excerpt) > 280 else "")
            cells[5].text = str(item["status"])

    document.add_heading("Resumen de validación automática", level=1)
    validation = payload.get("validation") or {}
    if validation:
        document.add_paragraph(f"Artefactos analizados: {validation.get('artifact_count', 0)}")
        document.add_paragraph(f"Estado de trazabilidad: {validation.get('traceability_status', 'no disponible')}")
        document.add_paragraph(f"Estado de calidad básica: {validation.get('quality_status', 'no disponible')}")
        document.add_paragraph(
            f"Posibles duplicados entre tipos: {len(validation.get('cross_type_duplicates', []))}"
        )
    else:
        document.add_paragraph("No existe un reporte de validación almacenado.")

    run = payload.get("generation_run") or {}
    configuration = (run.get("parameters") or {}).get("experimental_config") or {}
    if configuration:
        document.add_heading("Configuración de la ejecución", level=1)
        document.add_paragraph(
            f"Modelo LLM: {(configuration.get('llm') or {}).get('model', 'No registrado')}"
        )
        document.add_paragraph(
            f"Embeddings: {(configuration.get('embedding') or {}).get('model', 'No registrado')}"
        )
        retrieval = configuration.get("retrieval") or {}
        document.add_paragraph(
            f"Recuperación: {retrieval.get('method', 'No registrada')}; top-k: {retrieval.get('top_k', 'No registrado')}"
        )
        document.add_paragraph(f"Versión de prompts: {configuration.get('prompt_version', 'No registrada')}")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _priority_source_label(value: Any) -> str:
    return {
        "corpus": "Corpus",
        "usuario": "Decisión del usuario",
        "no_definida": "No definida",
        "legado": "Registro anterior sin procedencia",
    }.get(str(value or ""), "No registrada")


def _add_artifact_table(document: Any, item: dict[str, Any], artifact_type: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    header.text = f"{item['artifact_key']}. {item['title']}"
    if header.paragraphs and header.paragraphs[0].runs:
        header.paragraphs[0].runs[0].bold = True

    def add_row(label: str, value: Any) -> None:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        if cells[0].paragraphs and cells[0].paragraphs[0].runs:
            cells[0].paragraphs[0].runs[0].bold = True

    def list_value(values: list[Any] | None, empty: str) -> str:
        cleaned = [str(value).strip() for value in values or [] if str(value).strip()]
        return "\n".join(f"• {value}" for value in cleaned) if cleaned else empty

    add_row("Tipo", artifact_type)
    add_row("Redacción", item["description"])
    add_row("Prioridad", item["priority"])
    add_row("Procedencia de prioridad", _priority_source_label(item.get("priority_source")))
    add_row("Estado", item["status"])
    add_row("Versión", item["version"])
    add_row("Fuentes", ", ".join(item.get("source_fragments") or []) or "Sin evidencia asociada")
    add_row(
        "Artefactos relacionados",
        ", ".join(item.get("related_artifacts") or []) or "Sin relaciones registradas",
    )
    if item.get("rationale"):
        add_row("Justificación", item["rationale"])

    if artifact_type in {"RF", "RNF"}:
        add_row("Patrón de redacción", item.get("ears_pattern") or "No determinado")
        add_row(
            "Criterios de verificación",
            list_value(item.get("verification_criteria"), "Pendiente de definición"),
        )

    if artifact_type == "RNF":
        for label, key in (
            ("Categoría de calidad", "quality_category"),
            ("Métrica", "metric"),
            ("Unidad", "unit"),
            ("Umbral o valor objetivo", "target"),
            ("Método de verificación", "verification_method"),
        ):
            add_row(label, item.get(key) or "Pendiente de definición")

    if artifact_type == "HU":
        add_row(
            "Criterios de aceptación",
            list_value(item.get("acceptance_criteria"), "Pendiente de definición"),
        )

    warnings = (item.get("validation") or {}).get("warnings", [])
    if warnings:
        add_row("Observaciones automáticas", list_value(warnings, "Sin observaciones"))
